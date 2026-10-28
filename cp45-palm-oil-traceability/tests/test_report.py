"""The report builder must produce a well-formed DOCX with all 11 top-level sections, every figure, and no template leftovers."""
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

import _paths
from _paths import ROOT
import lib
import build_docx

EXPECTED_H1 = ["Executive Summary", "Chapter 1", "Chapter 2", "Chapter 3", "Chapter 4", "Chapter 5", "Chapter 6", "Chapter 7",
               "Appendix A", "Appendix B", "References"]


class ReportBuild(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        build_docx.build_doc([(1, "x", "0")], [("x", "x", "0")])         # pass 0 collects headings and captions
        cls.entries = [(l, t, "0") for l, t in lib.CTX.headings]
        cls.lof = [(a, b, "0") for a, b in lib.CTX.lof]
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "report.docx"
        build_docx.build_doc(cls.entries, cls.lof).save(cls.path)
        cls.doc = Document(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_top_level_sections(self):
        h1 = [t for l, t, _ in self.entries if l == 1]
        self.assertEqual(len(h1), len(EXPECTED_H1))
        for got, want in zip(h1, EXPECTED_H1):
            self.assertTrue(got.startswith(want), (got, want))

    def test_counts(self):
        self.assertEqual(sum(1 for a, _, _ in self.lof if a.startswith("Figure")), 18)
        self.assertEqual(sum(1 for a, _, _ in self.lof if a.startswith("Table")), 30)
        self.assertEqual(len(self.doc.inline_shapes), 18)

    def test_all_figure_files_exist(self):
        self.assertEqual(len(list((ROOT / "figures").glob("f*.png"))), 18)

    def test_no_template_leftovers(self):
        text = "\n".join(p.text for p in self.doc.paragraphs)
        for t in self.doc.tables:
            if len(t.columns) > 1:                                      # single-cell tables are the code listings in Appendix B
                text += "\n" + "\n".join(c.text for r in t.rows for c in r.cells)
        self.assertNotRegex(text, r"\bnan\b|\*\*|\{[a-z_]+\}|Traceback")
        for placeholder in ("Team member names and student IDs", "Faculty supervisor"):
            self.assertIn(placeholder, text)                            # cover placeholders are still there for the team to fill

    def test_honesty_statements_present(self):
        text = "\n".join(p.text for p in self.doc.paragraphs)
        self.assertIn("not estimated", text)
        self.assertIn("ASSUMPTION", text)


if __name__ == "__main__":
    unittest.main()
