"""Checks on the shipped evidence workbook, and (if LibreOffice is installed) on a fresh rebuild of it."""
import json
import math
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

import _paths
from _paths import DERIVED, WORKBOOK, ROOT

SHEETS = ['Start_Here', 'Summary', 'Assumptions', 'ROI_ULVR', 'ROI_NESN', 'ROI_WIL', 'MC_Results', 'Sensitivity',
          'Risk_Index', 'Supply_Chain', 'News_Events', 'ESG', 'Fundamentals', 'Market_Data', 'Price_Analytics', 'Next_Steps']
ERRORS = {"#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}


class ShippedWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.vals = openpyxl.load_workbook(WORKBOOK, data_only=True)
        cls.form = openpyxl.load_workbook(WORKBOOK)

    def test_sheets(self):
        self.assertEqual(self.vals.sheetnames, SHEETS)

    def test_no_formula_errors(self):
        bad = [(ws.title, c.coordinate, c.value) for ws in self.vals for row in ws.iter_rows() for c in row if c.value in ERRORS]
        self.assertEqual(bad, [])

    def test_formulas_have_cached_values(self):
        missing = []
        for ws in self.form:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and self.vals[ws.title][c.coordinate].value is None:
                        if '""' not in c.value:                         # IF(...,"") formulas legitimately return a blank (no price print that day)
                            missing.append((ws.title, c.coordinate))
        self.assertEqual(missing[:5], [])

    def test_spoei_scores_agree_with_derived_numbers(self):
        ws = self.vals["Risk_Index"]
        with open(DERIVED / "nums.json") as fh:
            nums = json.load(fh)
        for row, key in ((18, "base"), (19, "equal"), (20, "env")):
            got = [ws.cell(row, c).value for c in (8, 9, 10)]
            for a, b in zip(got, nums["score"][key]):
                self.assertAlmostEqual(a, b, places=9)

    def test_assumption_cells_are_flagged(self):
        text = " ".join(str(c.value) for row in self.vals["Assumptions"].iter_rows() for c in row if c.value)
        self.assertIn("ASSUMPTION", text)


@unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice not installed")
class Rebuild(unittest.TestCase):
    def test_rebuild_matches_shipped_workbook(self):
        import build_workbook
        with tempfile.TemporaryDirectory() as td:
            out = build_workbook.build(Path(td) / "rebuilt.xlsx")
            self.assertTrue(build_workbook.recalc(out))
            a = openpyxl.load_workbook(WORKBOOK, data_only=True)
            b = openpyxl.load_workbook(out, data_only=True)
            self.assertEqual(a.sheetnames, b.sheetnames)
            diffs = []
            for name in a.sheetnames:
                for row in a[name].iter_rows():
                    for c in row:
                        x, y = c.value, b[name][c.coordinate].value
                        same = math.isclose(x, y, rel_tol=1e-9, abs_tol=1e-9) if isinstance(x, (int, float)) and isinstance(y, (int, float)) else x == y
                        if not same:
                            diffs.append((name, c.coordinate, x, y))
            self.assertEqual(diffs[:5], [])


if __name__ == "__main__":
    unittest.main()
