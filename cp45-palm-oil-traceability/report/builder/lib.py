import re, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

import json, pathlib
ROOT = pathlib.Path(__file__).resolve().parents[2]
FIG = str(ROOT / "figures") + "/"
DER = ROOT / "data" / "derived"

def J(name):
    """Load a JSON file from data/derived/."""
    with open(DER / name) as fh:
        return json.load(fh)
NAVY = RGBColor(0x14, 0x2B, 0x4D); TEAL = RGBColor(0x0B, 0x6E, 0x8A); GREY = RGBColor(0x55, 0x5B, 0x66); INKC = RGBColor(0x1F, 0x29, 0x37)
FONT = "Calibri"; TEXT_W = 16.4  # cm

class Ctx:
    def __init__(self):
        self.fig = {}; self.tab = {}; self.lof = []; self.headings = []; self.chap = 0
CTX = Ctx()

def set_font(run, size=None, bold=None, italic=None, color=None, name=FONT, mono=False):
    n = "Courier New" if mono else name
    run.font.name = n; rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), n)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color

def shade(cell, hex_):
    tcPr = cell._element.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")): tcPr.remove(e)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)

def cell_borders(cell, **kw):
    tcPr = cell._element.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kw:
            v = kw[edge]; e = b.find(qn("w:" + edge))
            if e is None: e = OxmlElement("w:" + edge); b.append(e)
            if v is None: e.set(qn("w:val"), "nil")
            else:
                sz, col = v; e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), col)

def cell_margins(cell, top=50, bottom=50, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr(); m = tcPr.find(qn("w:tcMar"))
    if m is None: m = OxmlElement("w:tcMar"); tcPr.append(m)
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = m.find(qn("w:" + k))
        if e is None: e = OxmlElement("w:" + k); m.append(e)
        e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")

def para_border_bottom(p, color="142B4D", sz=10, space=4):
    pPr = p._element.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color); pb.append(b); pPr.append(pb)

def pfmt(p, before=0, after=6, line=1.12, keep_next=False, align=None, left=None, first=None, keep_lines=False):
    f = p.paragraph_format; f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = line
    f.keep_with_next = keep_next; f.keep_together = keep_lines
    if align is not None: p.alignment = align
    if left is not None: f.left_indent = Cm(left)
    if first is not None: f.first_line_indent = Cm(first)

_md = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`.+?`)")
def runs(p, text, size=10.5, color=INKC, bold=False, italic=False):
    for part in _md.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"): r = p.add_run(part[2:-2]); set_font(r, size, True, italic, color)
        elif part.startswith("`") and part.endswith("`"): r = p.add_run(part[1:-1]); set_font(r, size - 1, bold, italic, color, mono=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: r = p.add_run(part[1:-1]); set_font(r, size, bold, True, color)
        else: r = p.add_run(part); set_font(r, size, bold, italic, color)

def setup(doc):
    s = doc.sections[0]; s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.3); s.top_margin = Cm(2.3); s.bottom_margin = Cm(2.1); s.header_distance = Cm(1.1); s.footer_distance = Cm(1.0)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name, size, col, before, after in (("Heading 1", 21, NAVY, 0, 10), ("Heading 2", 13.5, NAVY, 14, 5), ("Heading 3", 11, TEAL, 10, 3)):
        h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = col; h.font.italic = False
        rf = h.element.rPr.find(qn("w:rFonts")) if h.element.rPr is not None else None
        if rf is not None:
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True

def field(run_parent_p, instr, cached="1", size=8.5, color=GREY):
    def mk(t):
        r = run_parent_p.add_run(); set_font(r, size, color=color); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._element.append(e); return r
    mk("begin"); r = run_parent_p.add_run(); set_font(r, size, color=color); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._element.append(it)
    mk("separate"); r = run_parent_p.add_run(cached); set_font(r, size, color=color); mk("end")

def header_footer(doc, left_text):
    s = doc.sections[0]; s.different_first_page_header_footer = True
    hp = s.header.paragraphs[0]; hp.text = ""; runs(hp, left_text, 8.5, GREY); para_border_bottom(hp, "C9CED6", 4, 3); pfmt(hp, after=0)
    fp = s.footer.paragraphs[0]; fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    runs(fp, "Page ", 8.5, GREY); field(fp, "PAGE", "2"); pfmt(fp, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)

def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); pfmt(p, after=0, line=1.0)

def h1(doc, text, chap=None, new_page=True, toc=True):
    if new_page: 
        p0 = doc.add_paragraph(); p0.paragraph_format.page_break_before = False
        p0._element.getparent().remove(p0._element)
    p = doc.add_paragraph(style="Heading 1"); p.paragraph_format.page_break_before = new_page
    r = p.add_run(text); set_font(r, 21, True, False, NAVY); para_border_bottom(p, "142B4D", 12, 6); pfmt(p, 0, 10, 1.0, True)
    if toc: CTX.headings.append((1, text))
    return p

def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2"); r = p.add_run(text); set_font(r, 13.5, True, False, NAVY); pfmt(p, 14, 5, 1.0, True); CTX.headings.append((2, text)); return p

def h3(doc, text):
    p = doc.add_paragraph(style="Heading 3"); r = p.add_run(text); set_font(r, 11, True, False, TEAL); pfmt(p, 9, 3, 1.0, True); return p

def para(doc, text, size=10.5, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=False, italic=False, color=INKC, left=None):
    p = doc.add_paragraph(); runs(p, text, size, color, italic=italic); pfmt(p, 0, after, 1.12, keep_next, align, left); return p

def bullets(doc, items, size=10.5, num=False, after=2, left=0.6):
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph(); pPr = p._element.get_or_add_pPr()
        p.paragraph_format.left_indent = Cm(left); p.paragraph_format.first_line_indent = Cm(-0.45)
        lead = f"{i}." if num else "•"
        tab = OxmlElement("w:tabs"); t_ = OxmlElement("w:tab"); t_.set(qn("w:val"), "left"); t_.set(qn("w:pos"), str(int(left * 567))); tab.append(t_); pPr.append(tab)
        r = p.add_run(lead + "\t"); set_font(r, size, color=TEAL, bold=True); runs(p, t, size)
        pfmt(p, 0, after, 1.1, False, WD_ALIGN_PARAGRAPH.LEFT, left, -0.45)
    if items: doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

def callout(doc, title, text, kind="key"):
    pal = {"key": ("E8F1FA", "0072B2"), "warn": ("FFF6E0", "C27C0E"), "ok": ("E9F6EE", "2E8B57"), "todo": ("FFF9C4", "B7950B")}[kind]
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    c = t.rows[0].cells[0]; c.width = Cm(TEXT_W); shade(c, pal[0]); cell_margins(c, 90, 90, 160, 140)
    cell_borders(c, left=(24, pal[1]), top=None, bottom=None, right=None)
    p = c.paragraphs[0]; runs(p, f"**{title}**", 10, RGBColor.from_string(pal[1])); pfmt(p, 0, 2, 1.1, True, WD_ALIGN_PARAGRAPH.LEFT)
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = c.add_paragraph(); runs(p, line, 10); pfmt(p, 0, 2 if i < len(text if isinstance(text, list) else [text]) - 1 else 0, 1.1, False, WD_ALIGN_PARAGRAPH.LEFT)
    sp = doc.add_paragraph(); pfmt(sp, 0, 4, 1.0)
    for r_ in sp.runs: r_.font.size = Pt(4)
    return t

def _tbl_props(t, widths):
    tbl = t._tbl; tblPr = tbl.tblPr
    for e in tblPr.findall(qn("w:tblLayout")): tblPr.remove(e)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths) * 567))); tw.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))

def table(doc, header, rows, widths, size=8.5, align=None, zebra=True, bold_first_col=False, hdr_fill="142B4D", total_row=False, highlight=None, pad=40):
    """rows: list of lists of str. align: list of 'l'/'r'/'c' per column. highlight: dict {(r,c): hex} or {r: hex}."""
    n = len(header); t = doc.add_table(rows=1 + len(rows), cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    _tbl_props(t, widths); align = align or ["l"] + ["r"] * (n - 1)
    AL = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
    for j, htxt in enumerate(header):
        c = t.rows[0].cells[j]; c.width = Cm(widths[j]); shade(c, hdr_fill); cell_margins(c, 60, 60, 80, 80); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; runs(p, htxt, size, RGBColor(255, 255, 255), bold=True); pfmt(p, 0, 0, 1.05, True, AL["l" if j == 0 else ("c" if align[j] == "r" else align[j])])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, row in enumerate(rows, 1):
        last = total_row and i == len(rows)
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); t.rows[i]._tr.get_or_add_trPr().append(cs)
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.width = Cm(widths[j]); cell_margins(c, pad, pad, 80, 80); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill = "F3F5F8" if (zebra and i % 2 == 0) else "FFFFFF"
            if last: fill = "E3E8EF"
            if highlight:
                if (i - 1, j) in highlight: fill = highlight[(i - 1, j)]
                elif (i - 1) in highlight: fill = highlight[i - 1]
            shade(c, fill); cell_borders(c, bottom=(4, "D5DAE1"), top=None, left=None, right=None)
            p = c.paragraphs[0]; runs(p, str(v), size, bold=(bold_first_col and j == 0) or last); pfmt(p, 0, 0, 1.05, False, AL[align[j]])
    sp = doc.add_paragraph(); pfmt(sp, 0, 2, 1.0)
    return t

def caption(doc, kind, chap, text, above=True):
    key = (kind, chap); CTX.fig[key] = CTX.fig.get(key, 0) + 1; n = CTX.fig[key]; label = f"{kind} {chap}.{n}"
    p = doc.add_paragraph(); runs(p, f"**{label} — {text}**", 9.5, NAVY); pfmt(p, 6 if above else 2, 3 if above else 8, 1.05, above, WD_ALIGN_PARAGRAPH.LEFT)
    CTX.lof.append((label, text)); return label

def figure(doc, chap, fname, title, source, width=TEXT_W, h=None):
    label = caption(doc, "Figure", chap, title, above=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; pfmt(p, 0, 2, 1.0, True, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(); r.add_picture(FIG + fname, width=Cm(width))
    s = doc.add_paragraph(); runs(s, source, 8.5, GREY, italic=True); pfmt(s, 0, 9, 1.05, False, WD_ALIGN_PARAGRAPH.LEFT)
    return label

def table_cap(doc, chap, title):
    return caption(doc, "Table", chap, title, above=True)

def note(doc, text, after=8):
    p = doc.add_paragraph(); runs(p, text, 8.5, GREY, italic=True); pfmt(p, 0, after, 1.05, False, WD_ALIGN_PARAGRAPH.LEFT); return p

def code(doc, text, size=7):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; _tbl_props(t, [TEXT_W])
    c = t.rows[0].cells[0]; c.width = Cm(TEXT_W); shade(c, "F6F7F9"); cell_margins(c, 60, 60, 120, 100); cell_borders(c, left=(18, "9CA3AF"), top=None, bottom=None, right=None)
    lines = text.rstrip("\n").split("\n")
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        r = p.add_run(ln if ln else " "); set_font(r, size, color=RGBColor(0x1F, 0x29, 0x37), mono=True); pfmt(p, 0, 0, 1.0, False, WD_ALIGN_PARAGRAPH.LEFT)
    sp = doc.add_paragraph(); pfmt(sp, 0, 4, 1.0)
