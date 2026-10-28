from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

F = "Arial"
NAVY = "1F3864"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color=NAVY)
f_sub = Font(name=F, size=10, italic=True, color="595959")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_in = Font(name=F, size=10, color="0000FF")      # hardcoded input / transcribed data
f_link = Font(name=F, size=10, color="008000")    # cross-sheet link
f_calc = Font(name=F, size=10, color="000000")
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_sec = PatternFill("solid", fgColor="D9E2F3")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_note = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

NUM = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = f_title
    if sub:
        ws["A2"] = sub; ws["A2"].font = f_sub
    ws.sheet_view.showGridLines = False

def hdr(ws, row, values, col=1, fill=fill_hdr, font=f_hdr, height=None):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = font; c.fill = fill; c.alignment = center; c.border = box
    if height: ws.row_dimensions[row].height = height

def sec(ws, row, text, ncols=8):
    for i in range(ncols):
        ws.cell(row=row, column=1 + i).fill = fill_sec
    ws.cell(row=row, column=1, value=text).font = f_bold

def put(ws, ref, value, font=None, fmt=None, fill=None, align=None, border=True):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    c.value = value
    c.font = font or f_norm
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = box
    return c

def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v

def note(ws, ref, text, width_cols=None):
    c = ws[ref]; c.value = text; c.font = f_sub; c.alignment = wrap
