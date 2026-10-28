from wb_common import *
import datetime as dt
import wb_a, wb_b1

D = dt.datetime
NEWS = [  # date, co, event, category, sentiment, severity(0-5), n_headlines, screen filter
 (D(2026,10,13), "WIL", "Wilmar tops global children's-rights benchmark for 5th consecutive year", "ESG – positive", 1, 0, 1, "Wilmar + Palm Oil"),
 (D(2026,9,11), "WIL", "Pakistan regulator probe vs Unity on Wilmar's complaint (Wilmar is complainant)", "Regulatory – non-adverse", 0, 0, 3, "Wilmar + Palm Oil"),
 (D(2026,9,7), "WIL", "ABF–Wilmar JV control of Hebei Mauri cleared in EU", "M&A – positive", 1, 0, 1, "Wilmar + Palm Oil"),
 (D(2026,9,24), "WIL", "Palm-oil prices capped for now; war risk, El Niño 2027 gains (BI)", "Commodity / price", 0, 0, 3, "Wilmar + Palm Oil"),
 (D(2026,8,12), "WIL", "Shares fall most in two months on earnings miss; 1H net income $608.9M vs $594.9M", "Earnings", -1, 1, 3, "Wilmar + Palm Oil"),
 (D(2026,5,28), "WIL", "Indonesia export under-invoicing / transfer-pricing probe of palm firms; Wilmar shares fall most in six years", "Regulatory / legal", -1, 5, 10, "Wilmar + ESG"),
 (D(2026,5,8), "WIL", "Russian government to seize Wilmar stake in indirect unit amid corruption allegations", "Legal / geopolitical", -1, 4, 1, "Wilmar + ESG"),
 (D(2026,4,8), "WIL", "Glass Lewis controversy alert on Wilmar", "Governance", -1, 2, 1, "Wilmar + ESG"),
 (D(2026,3,3), "WIL", "Indonesian court jails Wilmar unit executive (6 yrs) for bribery; shares fall", "Legal / corruption", -1, 4, 2, "Wilmar + ESG"),
 (D(2026,1,4), "WIL", "Indonesia's 'Switzerland-sized' land seizures tighten grip on plantations", "Land / deforestation", -1, 3, 1, "Wilmar + ESG"),
 (D(2025,7,15), "NESN", "Musim Mas, Nestlé and AAK renew palm partnership on environmental/social issues", "Mitigation – positive", 1, 0, 2, "Nestlé + Palm Oil"),
 (D(2026,6,16), "NESN", "Nestlé India higher as palm-oil prices decline, easing supply chain", "Commodity / price", 0, 0, 1, "Nestlé + Palm Oil"),
 (D(2026,4,9), "NESN", "Sell-side cost trackers flag building 2027 COGS-inflation risk", "Commodity / price", -1, 1, 4, "Nestlé + Palm Oil"),
 (D(2026,9,16), "ULVR", "UK CMA probe into McCormick's acquisition of Unilever Foods", "Antitrust (not supply chain)", -1, 2, 3, "Unilever + ESG"),
 (D(2026,9,3), "ULVR", "Unilever plant halts some production over quality-control failures", "Operational / quality", -1, 3, 1, "Unilever + ESG"),
 (D(2026,9,15), "ULVR", "Ben & Jerry's: 'toxic culture' allegations, board changes", "Governance / social", -1, 2, 3, "Unilever + ESG"),
 (D(2026,7,22), "ULVR", "Brand owners miss recycled-plastic packaging targets (BNEF)", "ESG – packaging", -1, 2, 1, "Unilever + ESG"),
 (D(2026,7,7), "ULVR", "Terry Smith accuses Unilever of misleading him over $66bn food deal", "Governance / investor", -1, 2, 2, "Unilever + ESG"),
 (D(2026,6,21), "ULVR", "Sunlight dishwashing liquid ordered to change 'misleading' claim", "Greenwashing / claims", -1, 3, 1, "Unilever + ESG"),
 (D(2026,5,15), "ULVR", "Turkey antitrust probe into Unilever and Magnum Dondurma", "Antitrust (not supply chain)", -1, 2, 1, "Unilever + ESG"),
 (D(2024,10,15), "ULVR", "Billionaire clan fuelled by green debt faces deforestation claims (context)", "Land / deforestation", -1, 3, 1, "Unilever + Palm Oil + deforestation"),
 (D(2023,2,15), "ULVR", "Indonesia tames rainforest destruction: crackdowns, boycotts and activism (WSJ, context)", "Land / deforestation", 0, 0, 1, "Unilever + Palm Oil + deforestation"),
]
NEWSROW = {}
def build_news(wb):
    ws = wb.create_sheet("News_Events")
    title(ws, "News & controversy log — hand-coded from Bloomberg CN screens",
          "One row per event cluster (many headlines on the same event count once). Antitrust events are logged but excluded from the adverse score (not supply-chain related). Category, sentiment and severity are MY judgement from headlines — re-code with your own rubric.")
    put(ws, "A4", "As-of date", f_bold); put(ws, "B4", D(2026,10,15), f_in, "dd-mmm-yy", fill=fill_key)
    put(ws, "C4", "12M window start", f_bold); put(ws, "D4", "=EDATE(B4,-12)", f_calc, "dd-mmm-yy")
    hdr(ws, 6, ["Company", "Adverse events (12M, excl. antitrust)", "Severity-weighted (12M)", "Positive/mitigation events (12M)", "Headlines in log (12M)", "Screen filter used"], height=44)
    a, b = 14, 14 + len(NEWS) - 1
    flt = {"WIL": "ESG + Palm Oil", "NESN": "Palm Oil ONLY", "ULVR": "ESG (+ Palm Oil/deforestation, older)"}
    for i, co in enumerate(("ULVR", "NESN", "WIL")):
        r = 7 + i
        put(ws, (r, 1), co, f_bold)
        put(ws, (r, 2), f'=COUNTIFS($B${a}:$B${b},A{r},$H${a}:$H${b},TRUE,$E${a}:$E${b},-1,$D${a}:$D${b},"<>Antitrust*")', f_calc, '0')
        put(ws, (r, 3), f'=SUMIFS($I${a}:$I${b},$B${a}:$B${b},A{r})', f_calc, '0', fill=fill_key)
        put(ws, (r, 4), f'=COUNTIFS($B${a}:$B${b},A{r},$H${a}:$H${b},TRUE,$E${a}:$E${b},1)', f_calc, '0')
        put(ws, (r, 5), f'=SUMIFS($G${a}:$G${b},$B${a}:$B${b},A{r},$H${a}:$H${b},TRUE)', f_calc, '0')
        put(ws, (r, 6), flt[co], f_in)
        NEWSROW[co] = r
    note(ws, "A11", "WARNING — not like-for-like: the three companies were screened with different search filters (Nestlé: 'Palm Oil' only, which returns mostly sell-side cost notes). "
         "Re-run all three with the SAME query, e.g. company AND (palm oil OR deforestation OR 'supply chain') with ESG/controversy tags, before relying on the news score. Weight kept low in the index for this reason.")
    ws.merge_cells("A11:J11"); ws.row_dimensions[11].height = 44
    hdr(ws, 13, ["Date", "Co.", "Event cluster", "Category", "Sentiment (-1/0/+1)", "Severity 0–5", "# headlines", "In 12M window?", "Adverse severity counted", "Screen filter"], height=32)
    for i, (d, co, ev, cat, sen, sev, n, flt_) in enumerate(NEWS):
        r = a + i
        put(ws, (r, 1), d, f_in, "dd-mmm-yy"); put(ws, (r, 2), co, f_in); put(ws, (r, 3), ev, f_in, align=wrap)
        put(ws, (r, 4), cat, f_in); put(ws, (r, 5), sen, f_in, '0'); put(ws, (r, 6), sev, f_in, '0'); put(ws, (r, 7), n, f_in, '0')
        put(ws, (r, 8), f"=AND(A{r}>=$D$4,A{r}<=$B$4)", f_calc); put(ws, (r, 9), f'=IF(AND(H{r},E{r}<0,LEFT(D{r},9)<>"Antitrust"),F{r},0)', f_calc, '0')
        put(ws, (r, 10), flt_, f_in)
    widths(ws, {"A": 14, "B": 18, "C": 72, "D": 26, "E": 13, "F": 12, "G": 11, "H": 12, "I": 13, "J": 32})
    ws.freeze_panes = "A14"
    return ws

RI = {}
def build_risk(wb):
    ws = wb.create_sheet("Risk_Index")
    title(ws, "Composite palm-oil traceability risk index (0 = low risk, 100 = high risk) — real indicators",
          "Replaces the draft's synthetic 50-supplier index. Each indicator is scaled 0–100 between a 'no-risk' value (floor) and a 'max-risk' value (cap). Edit blue cells; scores update.")
    hdr(ws, 4, ["Indicator", "Source", "No-risk value", "Max-risk value", "Unilever (raw)", "Nestlé (raw)", "Wilmar (raw)",
                "Unilever (0–100)", "Nestlé (0–100)", "Wilmar (0–100)", "Weight: Base", "Weight: Equal", "Weight: Env/Social-heavy"], height=44)
    fx = {"ULVR": "Assumptions!$B$11", "NESN": "Assumptions!$B$12", "WIL": "1"}
    E, Fd, S = wb_a.ESGROW, wb_a.ROW, wb_b1.SC
    def esg(co, key): return f"ESG!J{E[co][key]}"
    def revusd(co): return f"Fundamentals!H{Fd[co]['rev']}*{fx[co]}/1000"
    inds = [
     ("BESG ESG score gap (10 − score, 2024)", "ESG", 0, 10, lambda co: f"=10-{esg(co,'besg')}", NUM2, (15, 10, 15)),
     ("BESG Social pillar gap (10 − score, 2024)", "ESG", 0, 10, lambda co: f"=10-{esg(co,'soc')}", NUM2, (15, 10, 20)),
     ("Scope 3 intensity (kt CO2e per USD bn revenue, 2024)", "ESG, Fundamentals, FX", 0, 3000, lambda co: f"={esg(co,'s3')}/({revusd(co)})", NUM2, (15, 10, 20)),
     ("Suppliers with business-conduct risk, %", "Supply_Chain (SPLC)", 0, 0.40, lambda co: f"=Supply_Chain!{ {'WIL':'B','NESN':'C','ULVR':'D'}[co] }{S['conduct']}", PCT, (10, 10, 10)),
     ("Suppliers with high GHG vs peers, %", "Supply_Chain (SPLC)", 0, 0.25, lambda co: f"=Supply_Chain!{ {'WIL':'B','NESN':'C','ULVR':'D'}[co] }{S['ghg']}", PCT, (5, 10, 10)),
     ("Suppliers that are ESG-score laggards, %", "Supply_Chain (SPLC)", 0, 0.10, lambda co: f"=Supply_Chain!{ {'WIL':'B','NESN':'C','ULVR':'D'}[co] }{S['lag']}", PCT, (5, 10, 5)),
     ("Sanctioned + distressed suppliers, %", "Supply_Chain (SPLC)", 0, 0.05, lambda co: f"=Supply_Chain!{ {'WIL':'B','NESN':'C','ULVR':'D'}[co] }{S['sanc_dist']}", PCT, (10, 10, 5)),
     ("Adverse news severity, 12M (see warning on News_Events)", "News_Events", 0, 20, lambda co: f"=News_Events!C{NEWSROW[co]}", NUM0, (10, 10, 5)),
     ("Equity volatility, 90-day (% ann.)", "Market_Data", 10, 40, lambda co: f"=Market_Data!{ {'WIL':'B','NESN':'C','ULVR':'D'}[co] }{wb_a.SNAPROW['90-day volatility (Bloomberg chart, % ann.)']}", NUM, (5, 10, 0)),
     ("Interest cover, 2025 (OI / interest) — lower = riskier", "Fundamentals", 12, 2, lambda co: f"=Fundamentals!I{Fd[co]['cov']}", MULT, (10, 10, 10)),
    ]
    r0 = 5
    for i, (lab, src, fl, cp, fn, fmt, wts) in enumerate(inds):
        r = r0 + i
        put(ws, (r, 1), lab); put(ws, (r, 2), src, f_sub)
        put(ws, (r, 3), fl, f_in, fmt); put(ws, (r, 4), cp, f_in, fmt)
        for j, co in enumerate(("ULVR", "NESN", "WIL")):
            put(ws, (r, 5 + j), fn(co), f_link, fmt)
            raw = f"{L(5+j)}{r}"
            put(ws, (r, 8 + j), f"=MAX(0,MIN(100,({raw}-$C{r})/($D{r}-$C{r})*100))", f_calc, NUM)
        for k, w in enumerate(wts):
            put(ws, (r, 11 + k), w / 100, f_in, '0%')
    r1 = r0 + len(inds) - 1
    r = r1 + 1
    put(ws, (r, 1), "Weight check (must be 100%)", f_bold)
    for k in range(3):
        put(ws, (r, 11 + k), f"=SUM({L(11+k)}{r0}:{L(11+k)}{r1})", f_calc, '0%')
    r += 2
    hdr(ws, r, ["Composite score", "", "", "", "", "", "", "Unilever", "Nestlé", "Wilmar"]); r += 1
    RI["score_rows"] = {}
    for k, nm in enumerate(["Base weights", "Equal weights", "Env/Social-heavy"]):
        put(ws, (r, 1), nm, f_bold)
        for j in range(3):
            put(ws, (r, 8 + j), f"=SUMPRODUCT({L(8+j)}${r0}:{L(8+j)}${r1},${L(11+k)}${r0}:${L(11+k)}${r1})", f_calc, NUM, fill=fill_key if k == 0 else None)
        RI["score_rows"][nm] = r; r += 1
    base_r = RI["score_rows"]["Base weights"]
    for k, nm in enumerate(["Base weights", "Equal weights", "Env/Social-heavy"]):
        rr = RI["score_rows"][nm]
        put(ws, (r, 1), f"Rank (1 = highest risk) — {nm}")
        for j in range(3):
            put(ws, (r, 8 + j), f"=RANK({L(8+j)}{rr},$H{rr}:$J{rr},0)", f_calc, '0')
        r += 1
    put(ws, (r, 1), "Ranking stable across the three weight sets?", f_bold)
    rk = r - 3
    put(ws, (r, 8), f'=IF(AND(H{rk}=H{rk+1},H{rk}=H{rk+2},I{rk}=I{rk+1},I{rk}=I{rk+2},J{rk}=J{rk+1},J{rk}=J{rk+2}),"Yes","No — weights matter")', f_calc)
    RI["stable"] = r
    r += 2
    note(ws, f"A{r}", "How to read: with only three firms the index is a RELATIVE ranking and a worked example of the method — not a validated risk model. "
         "To make it publishable: (1) extend to 10–15 firms (add IOI, Musim Mas, AAK, Cargill, P&G, Mondelez, Colgate); (2) score Wilmar's 30 suppliers at node level; "
         "(3) calibrate weights against an outcome (e.g. forward 3-month drawdown or controversy count after each event) instead of choosing them by judgement. "
         "Unilever's Social-pillar gap dominates its score; Wilmar's is driven by interest cover, supplier conduct risk and news severity.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13); ws.row_dimensions[r].height = 62
    widths(ws, {"A": 56, "B": 22, "C": 11, "D": 11, **{L(i): 13 for i in range(5, 14)}})
    ws.freeze_panes = "B5"
    from openpyxl.chart import BarChart, Reference
    ch = BarChart(); ch.type = "col"; ch.title = "Composite risk score (base weights)"; ch.height = 7; ch.width = 14
    ch.add_data(Reference(ws, min_col=8, max_col=10, min_row=base_r), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=8, max_col=10, min_row=base_r - 1))
    ch.legend = None; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100
    ws.add_chart(ch, f"A{r+2}")
    return ws
