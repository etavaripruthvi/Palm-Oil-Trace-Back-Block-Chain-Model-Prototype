from wb_common import *
from openpyxl.worksheet.datavalidation import DataValidation
import wb_a

PROF = [("ULVR", "Unilever", "B"), ("NESN", "Nestlé", "C"), ("WIL", "Wilmar", "D")]
AS = {}   # registry of assumption cells

def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions — every judgement input for the ROI / NPV model lives here",
          "Blue = input you can change (yellow = most influential). Green = pulled from Fundamentals. These are PLANNING-GRADE PLACEHOLDERS, not Bloomberg facts: calibrate each one and cite the source in the Basis column.")
    sec(ws, 4, "A. Global inputs", 6)
    G = [(5, "Discount rate", 0.08, PCT, "Use each firm's WACC (Bloomberg WACC <GO>); 8% is a generic consumer-staples/agri placeholder.", True),
         (6, "Nominal growth of revenue, COGS and opex", 0.02, PCT, "Placeholder; replace with consensus growth (see Fundamentals 2026E/27E).", False),
         (7, "Realised fine, % of EU turnover when non-compliant", 0.02, PCT, "EUDR maximum is at least 4% of EU turnover; base case assumes half of that is realised. Verify against the final regulation text.", True),
         (8, "Price premium on re-sourced compliant volume", 0.06, PCT, "Placeholder; calibrate to RSPO / segregated-supply premiums.", False),
         (9, "Share of non-traceable EU-bound volume excluded / re-sourced", 0.30, PCT, "Placeholder; supplier-exclusion rate under a compliance-only sourcing policy.", False),
         (10, "Baseline traceable share today (c0)", 0.30, PCT, "Placeholder; use each firm's published certified / traceable-to-mill percentage.", True),
         (11, "FX: EUR→USD", 1.15, NUM2, "ILLUSTRATIVE — replace with Bloomberg FX on the data date (EURUSD Curncy).", False),
         (12, "FX: CHF→USD", 1.25, NUM2, "ILLUSTRATIVE — replace with Bloomberg FX on the data date (USDCHF Curncy, inverted).", False)]
    for r, lab, v, fmt, basis, key in G:
        put(ws, (r, 1), lab); put(ws, (r, 2), v, f_in, fmt, fill=fill_key if key else None); put(ws, (r, 6), basis, f_sub, align=wrap, border=False)
    put(ws, "A13", "Active scenario (pick: Base / Adverse / Severe)", f_bold)
    put(ws, "B13", "Base", f_in, fill=fill_key, align=center)
    dv = DataValidation(type="list", formula1='"Base,Adverse,Severe"', allow_blank=False); ws.add_data_validation(dv); dv.add("B13")
    put(ws, "F13", "Drives the three ROI sheets and the live numbers on Summary. Monte Carlo tables are static (run in Python).", f_sub, border=False)
    hdr(ws, 15, ["Scenario multipliers", "Regulatory prob. ×", "Controversy prob. ×", "Substitution premium ×", "CPO price level ×"], height=32)
    for i, (nm, a, b, c, d) in enumerate([("Base", 1, 1, 1, 1.0), ("Adverse", 2, 1.5, 1.5, 1.0), ("Severe", 4, 2.5, 2.5, 1.3)]):
        put(ws, (16 + i, 1), nm, f_bold)
        for j, v in enumerate((a, b, c, d)): put(ws, (16 + i, 2 + j), v, f_in, '0.0"x"')
    put(ws, "A19", "Active multipliers", f_bold)
    for j in range(4):
        c = L(2 + j)
        put(ws, (19, 2 + j), f"=INDEX({c}16:{c}18,MATCH($B$13,$A$16:$A$18,0))", f_calc, '0.0"x"')
    put(ws, "F16", "Severe = strict EUDR enforcement + palm price level +30% (supply squeeze).", f_sub, border=False)

    sec(ws, 21, "B. Firm profiles (USD m unless stated)", 6)
    hdr(ws, 22, ["Parameter", "Unilever", "Nestlé", "Wilmar", "", "Basis"], height=24)
    Fd = wb_a.ROW
    rows = [
     (23, "Revenue FY2025 (reporting ccy m)", lambda co, pc: f"=Fundamentals!I{Fd[co]['rev']}", NUM0, f_link, "Bloomberg FA (EUR / CHF / USD)"),
     (24, "COGS FY2025 (reporting ccy m)", lambda co, pc: f"=Fundamentals!I{Fd[co]['cogs']}", NUM0, f_link, "Bloomberg FA"),
     (25, "FX to USD", lambda co, pc: {"ULVR": "=$B$11", "NESN": "=$B$12", "WIL": 1}[co], NUM2, f_calc, "Wilmar reports in USD"),
     (26, "Revenue FY2025 (USD m)", lambda co, pc: f"={pc}23*{pc}25", NUM0, f_calc, ""),
     (27, "COGS FY2025 (USD m)", lambda co, pc: f"={pc}24*{pc}25", NUM0, f_calc, ""),
    ]
    for r, lab, fn, fmt, font, basis in rows:
        put(ws, (r, 1), lab)
        for co, nm, pc in PROF:
            put(ws, (r, "BCD".index(pc) + 2), fn(co, pc), font, fmt)
        put(ws, (r, 6), basis, f_sub, border=False)
    inputs = [
     (28, "a  Palm-linked share of revenue", (0.25, 0.15, 0.40), PCT, "ASSUMPTION — share of sales in products with palm-derived inputs (home/personal care, foods; for Wilmar: palm & laurics, consumer products)."),
     (29, "ps  Palm-derived spend as % of COGS", (0.030, 0.015, 0.300), PCT, "ASSUMPTION — calibrate from sustainability reports (tonnes palm × price) divided by COGS."),
     (30, "eu  Revenue share in EU-regulated markets", (0.25, 0.25, 0.08), PCT, "ASSUMPTION — Wilmar low (China 47% of revenue, see Supply_Chain); brand owners via EU retail."),
     (31, "p_reg  Annual probability of a regulatory / market-access event", (0.03, 0.03, 0.05), PCT, "ASSUMPTION — to calibrate from enforcement history; the Wilmar Indonesia probe (May-26) shows the order of magnitude is not zero."),
     (32, "p_ctrv  Annual probability of a controversy / boycott event", (0.10, 0.10, 0.15), PCT, "ASSUMPTION — calibrate from the News_Events frequency (extend to ≥5Y)."),
     (33, "rev_loss  Revenue lost on palm-linked sales in a controversy", (0.005, 0.005, 0.010), '0.0%', "ASSUMPTION — event-study the 2026 Wilmar probe (needs daily prices, see Next_Steps)."),
     (34, "Central database: capex (USD m)", (8.0, 7.0, 25.0), NUM, "ASSUMPTION — planning-grade; replace with vendor quotes / pilot budgets."),
     (35, "Central database: annual opex (USD m)", (3.0, 2.5, 10.0), NUM, "ASSUMPTION"),
     (36, "Permissioned blockchain: capex (USD m)", (14.0, 12.0, 40.0), NUM, "ASSUMPTION — consortium build, node onboarding, integration."),
     (37, "Permissioned blockchain: annual opex (USD m)", (4.5, 4.0, 14.0), NUM, "ASSUMPTION"),
    ]
    for r, lab, vals, fmt, basis in inputs:
        put(ws, (r, 1), lab)
        for j, v in enumerate(vals): put(ws, (r, 2 + j), v, f_in, fmt)
        put(ws, (r, 6), basis, f_sub, border=False)
    sec(ws, 39, "C. Option design", 6)
    hdr(ws, 40, ["Parameter", "Status quo", "Central database", "Permissioned blockchain", "", "Basis"], height=32)
    put(ws, "A41", "cmax  Steady-state traceable share"); put(ws, "B41", "=B10", f_calc, PCT); put(ws, "C41", 0.80, f_in, PCT); put(ws, "D41", 0.90, f_in, PCT)
    put(ws, "A42", "Ramp-up years to reach cmax"); put(ws, "B42", 1, f_in, '0.0'); put(ws, "C42", 3, f_in, '0.0'); put(ws, "D42", 4, f_in, '0.0')
    put(ws, "A43", "eff  Assurance value of coverage (0–1)"); put(ws, "B43", 0.60, f_in, '0.00'); put(ws, "C43", 0.80, f_in, '0.00', fill=fill_key); put(ws, "D43", 0.90, f_in, '0.00', fill=fill_key)
    put(ws, "F41", "ASSUMPTION — blockchain coverage slightly higher (shared ledger across traders/mills) but slower to coordinate.", f_sub, border=False)
    put(ws, "F43", "ASSUMPTION — KEY: how much extra credibility an immutable multi-party ledger earns with regulators/retailers vs a central database. This single gap decides RQ4.", f_sub, border=False)
    note(ws, "A45", "Mechanics (same in Python): coverage c(t)=c0+(cmax−c0)·min(1,t/ramp); mitigation m=c·eff. Expected annual loss = "
         "[p_reg·fine·eu·Revenue + p_ctrv·rev_loss·a·Revenue]·(1−m) + PalmSpend·eu·(1−c)·q·π. Benefit = PV(total cost, status quo) − PV(total cost, option). Year 0 = build (capex only); years 1–7 = operation.")
    ws.merge_cells("A45:F45"); ws.row_dimensions[45].height = 52
    widths(ws, {"A": 58, "B": 16, "C": 18, "D": 22, "E": 18, "F": 95})
    ws.freeze_panes = "B4"
    return ws

ROIOUT = {}
def build_roi(wb, co, nm, pc):
    ws = wb.create_sheet(f"ROI_{co}")
    title(ws, f"{nm} — traceability programme: status quo vs central database vs permissioned blockchain (USD m)",
          "Pure calculation sheet: all inputs are green links to 'Assumptions'. Change the scenario on Assumptions!B13. Year 0 = build year.")
    A = "Assumptions!"
    inp = [(4, "Revenue, USD m (FY2025)", f"={A}{pc}26", NUM0), (5, "COGS, USD m (FY2025)", f"={A}{pc}27", NUM0),
           (6, "a  palm-linked revenue share", f"={A}{pc}28", PCT), (7, "ps  palm spend / COGS", f"={A}{pc}29", PCT),
           (8, "eu  EU-regulated revenue share", f"={A}{pc}30", PCT),
           (9, "p_reg (scenario-adjusted)", f"={A}{pc}31*{A}$B$19", '0.0%'), (10, "p_ctrv (scenario-adjusted)", f"={A}{pc}32*{A}$C$19", '0.0%'),
           (11, "rev_loss", f"={A}{pc}33", '0.0%'), (12, "Realised fine rate", f"={A}$B$7", PCT),
           (13, "Substitution premium π (scenario-adjusted)", f"={A}$B$8*{A}$D$19", PCT), (14, "Share excluded / re-sourced q", f"={A}$B$9", PCT),
           (15, "Discount rate", f"={A}$B$5", PCT), (16, "Growth", f"={A}$B$6", PCT),
           (17, "CPO price level × (scenario)", f"={A}$E$19", '0.00"x"'), (18, "Baseline coverage c0", f"={A}$B$10", PCT)]
    for r, lab, f, fmt in inp:
        put(ws, (r, 1), lab); put(ws, (r, 2), f, f_link, fmt)
    put(ws, "D4", "Active scenario:", f_bold, border=False); put(ws, "E4", f"={A}B13", f_link, fill=fill_key)
    hdr(ws, 20, ["Option design", "cmax", "Ramp (yrs)", "eff", "Capex", "Opex / yr"])
    opts = [(21, "Status quo", "B", None, None), (22, "Central database", "C", 34, 35), (23, "Permissioned blockchain", "D", 36, 37)]
    for r, lab, oc, cx, ox in opts:
        put(ws, (r, 1), lab, f_bold)
        put(ws, (r, 2), f"={A}{oc}41", f_link, PCT); put(ws, (r, 3), f"={A}{oc}42", f_link, '0.0'); put(ws, (r, 4), f"={A}{oc}43", f_link, '0.00')
        put(ws, (r, 5), f"={A}{pc}{cx}" if cx else 0, f_link if cx else f_calc, NUM)
        put(ws, (r, 6), f"={A}{pc}{ox}" if ox else 0, f_link if ox else f_calc, NUM)
    hdr(ws, 26, ["Year"] + list(range(0, 8)), col=2)
    put(ws, "A26", "Year", f_hdr, fill=fill_hdr)
    cols = [L(3 + t) for t in range(8)]   # C..J
    lines = [(27, "Growth factor", lambda c, t: f"=(1+$B$16)^{c}$26", '0.000'),
             (28, "Revenue", lambda c, t: f"=$B$4*{c}27", NUM0), (29, "COGS", lambda c, t: f"=$B$5*{c}27", NUM0),
             (30, "Discount factor", lambda c, t: f"=1/(1+$B$15)^{c}$26", '0.000'),
             (31, "CPO price index", lambda c, t: "=1" if t == 0 else "=$B$17", '0.00'),
             (32, "Palm-derived spend", lambda c, t: f"={c}29*$B$7*{c}31", NUM0)]
    for r, lab, fn, fmt in lines:
        put(ws, (r, 1), lab)
        put(ws, (r, 2), None)
        for t, c in enumerate(cols): put(ws, (r, 3 + t), fn(c, t), f_calc, fmt)
    blk = {}
    for k, (orow, oname) in enumerate([(21, "Status quo"), (22, "Central database"), (23, "Permissioned blockchain")]):
        s = 34 + 11 * k
        sec(ws, s, f"Option: {oname}", 10)
        names = ["Traceable coverage", "Mitigation (coverage × eff)", "Regulatory loss (expected)", "Supplier substitution cost",
                 "Reputational / controversy loss (expected)", "Programme opex", "Programme capex", "Total cost (losses + programme)", "PV of total cost"]
        for i, nme in enumerate(names): put(ws, (s + 1 + i, 1), nme); put(ws, (s + 1 + i, 2), None)
        for t, c in enumerate(cols):
            cov, m, reg, sub, rep, opx, cpx, tot, pv = [s + 1 + i for i in range(9)]
            if t == 0:
                put(ws, (cov, 3), "=$B$18", f_calc, PCT); put(ws, (m, 3), f"=C{cov}*$D${orow}", f_calc, PCT)
                for rr in (reg, sub, rep, opx): put(ws, (rr, 3), 0, f_calc, NUM)
                put(ws, (cpx, 3), f"=$E${orow}", f_calc, NUM)
            else:
                put(ws, (cov, 3 + t), f"=$B$18+($B${orow}-$B$18)*MIN(1,{c}$26/$C${orow})", f_calc, PCT)
                put(ws, (m, 3 + t), f"={c}{cov}*$D${orow}", f_calc, PCT)
                put(ws, (reg, 3 + t), f"=$B$9*$B$12*$B$8*{c}$28*(1-{c}{m})", f_calc, NUM)
                put(ws, (sub, 3 + t), f"={c}$32*$B$8*(1-{c}{cov})*$B$14*$B$13", f_calc, NUM)
                put(ws, (rep, 3 + t), f"=$B$10*$B$11*$B$6*{c}$28*(1-{c}{m})", f_calc, NUM)
                put(ws, (opx, 3 + t), f"=$F${orow}*{c}$27", f_calc, NUM)
                put(ws, (cpx, 3 + t), 0, f_calc, NUM)
            put(ws, (tot, 3 + t), f"=SUM({c}{reg}:{c}{cpx})", f_bold, NUM)
            put(ws, (pv, 3 + t), f"={c}{tot}*{c}$30", f_calc, NUM)
        blk[orow] = dict(reg=s + 3, sub=s + 4, rep=s + 5, opx=s + 6, cpx=s + 7, tot=s + 8, pv=s + 9)
    # summary
    s = 34 + 33 + 1
    sec(ws, s, "Summary (PV at year 0, USD m)", 10)
    hdr(ws, s + 1, ["Option", "PV expected losses", "PV programme cost", "PV total cost", "Benefit vs status quo (NPV)", "PV avoided losses", "Benefit / programme cost", "Break-even p_reg multiplier", "Discounted payback (yr)"], height=44)
    D = "$C$30:$J$30"
    rr = {}
    for i, (orow, oname) in enumerate([(21, "Status quo"), (22, "Central database"), (23, "Permissioned blockchain")]):
        r = s + 2 + i; rr[orow] = r; b = blk[orow]
        put(ws, (r, 1), oname, f_bold)
        put(ws, (r, 2), f"=SUMPRODUCT((C{b['reg']}:J{b['reg']}+C{b['sub']}:J{b['sub']}+C{b['rep']}:J{b['rep']})*{D})", f_calc, NUM)
        put(ws, (r, 3), f"=SUMPRODUCT((C{b['opx']}:J{b['opx']}+C{b['cpx']}:J{b['cpx']})*{D})", f_calc, NUM)
        put(ws, (r, 4), f"=SUM(C{b['pv']}:J{b['pv']})", f_calc, NUM)
    q = rr[21]
    for orow in (22, 23):
        r = rr[orow]; b = blk[orow]; bs = blk[21]
        put(ws, (r, 5), f"=$D${q}-D{r}", f_bold, NUM, fill=fill_key)
        put(ws, (r, 6), f"=$B${q}-B{r}", f_calc, NUM)
        put(ws, (r, 7), f"=IF(C{r}=0,0,F{r}/C{r})", f_calc, MULT)
        put(ws, (r, 8), f"=IFERROR(MAX(0,1-E{r}/SUMPRODUCT((C{bs['reg']}:J{bs['reg']}-C{b['reg']}:J{b['reg']})*{D})),0)", f_calc, '0.00"x"')
    for c in range(5, 10): put(ws, (q, c), None)
    # payback helpers
    h = s + 7
    put(ws, (h, 1), "Cumulative discounted net benefit vs status quo", f_bold, border=False)
    for i, orow in enumerate((22, 23)):
        r = h + 1 + 2 * i; b = blk[orow]; bs = blk[21]
        put(ws, (r, 1), ["Central database — cumulative", "Permissioned blockchain — cumulative"][i]); put(ws, (r, 2), None)
        put(ws, (r + 1, 1), "  payback flag (year if cumulative ≥ 0, else 99)"); put(ws, (r + 1, 2), None)
        for t, c in enumerate(cols):
            prev = f"+{L(2+t)}{r}" if t > 0 else ""
            put(ws, (r, 3 + t), f"=({c}{bs['tot']}-{c}{b['tot']})*{c}$30{prev}", f_calc, NUM)
            put(ws, (r + 1, 3 + t), f"=IF({c}{r}>=0,{c}$26,99)", f_calc, '0')
        put(ws, (rr[orow], 9), f'=IF(MIN(C{r+1}:J{r+1})=99,"> 7",MIN(C{r+1}:J{r+1}))', f_calc, '0', align=center)
    r = h + 6
    put(ws, (r, 1), "Blockchain minus central database (NPV)", f_bold); put(ws, (r, 5), f"=E{rr[23]}-E{rr[22]}", f_bold, NUM, fill=fill_key)
    put(ws, (r + 1, 1), "Break-even p_reg multiplier: the factor on the regulatory-event probability at which the option's NPV is zero (0 = pays back even with no regulatory risk, because substitution and reputation benefits already cover costs).", f_sub, border=False)
    ROIOUT[co] = dict(DB=rr[22], BC=rr[23], diff=r)
    widths(ws, {"A": 46, "B": 16, **{L(i): 13 for i in range(3, 11)}})
    ws.freeze_panes = "B4"
    return ws
