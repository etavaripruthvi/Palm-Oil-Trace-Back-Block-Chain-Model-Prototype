from wb_common import *
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule
import numpy as np
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
import palm_oil_traceability_model as M

NAMES = {"ULVR": "Unilever", "NESN": "Nestlé", "WIL": "Wilmar"}
MCROW = {}
def build_mc(wb):
    ws = wb.create_sheet("MC_Results")
    title(ws, f"Monte Carlo results — {M.N_SIMS:,} draws per cell, seed {M.SEED} (STATIC values from palm_oil_traceability_model.py)",
          "NPV benefit vs status quo, USD m. Same random world is applied to every option (common random numbers). Re-run the script after changing assumptions; these cells do not update by themselves.")
    mc = M.mc_summary()
    hdr(ws, 4, ["Firm", "Regime", "Option", "Mean", "5th pct", "Median", "95th pct", "P(NPV > 0)", "P(blockchain beats database)", "Mean (blockchain − database)"], height=44)
    r = 5
    for _, x in mc.iterrows():
        put(ws, (r, 1), NAMES[x.profile]); put(ws, (r, 2), x.regime); put(ws, (r, 3), "Central database" if x.option == "DB" else "Permissioned blockchain")
        for j, k in enumerate(["mean", "p5", "median", "p95"]): put(ws, (r, 4 + j), float(x[k]), f_in, NUM)
        put(ws, (r, 8), float(x.p_positive), f_in, '0%')
        if x.option == "BC":
            put(ws, (r, 9), float(x.p_BC_beats_DB), f_in, '0%'); put(ws, (r, 10), float(x.mean_BC_minus_DB), f_in, NUM)
        else:
            put(ws, (r, 9), None); put(ws, (r, 10), None)
        MCROW[(x.profile, x.regime, x.option)] = r
        r += 1
    note(ws, f"A{r+1}", "Regimes: Base / Adverse / Severe use the multipliers on Assumptions; 'Mixed' draws a regime per simulation (25% benign EU enforcement, 50% base, 25% severe). "
         "Random inputs: event probabilities (log-normal), realised fine 0.5–4%, exclusion share, premium, discount rate 6–11%, baseline coverage, option coverage/ramp/effectiveness (triangular), cost overruns (log-normal), "
         f"and a CPO price path (GBM, zero drift, σ = {M.CPO_SIGMA:.0%} estimated from only ~8 weeks of KOZ6 prints — re-estimate on 5Y data).")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=10); ws.row_dimensions[r + 1].height = 64
    # histogram for Unilever mixed
    b = M.monte_carlo("ULVR", "Mixed")
    lo, hi = np.percentile(np.r_[b["DB"], b["BC"]], [1, 99])
    edges = np.linspace(lo, hi, 21)
    cdb, _ = np.histogram(np.clip(b["DB"], lo, hi), edges); cbc, _ = np.histogram(np.clip(b["BC"], lo, hi), edges)
    h0 = r + 4
    sec(ws, h0, "Distribution of NPV benefit — Unilever, mixed regimes (share of simulations per bin; tails clipped at 1st/99th pct)", 10)
    hdr(ws, h0 + 1, ["Bin mid-point (USD m)", "Central database", "Permissioned blockchain"])
    for i in range(20):
        put(ws, (h0 + 2 + i, 1), round(float((edges[i] + edges[i + 1]) / 2), 1), f_in, NUM)
        put(ws, (h0 + 2 + i, 2), float(cdb[i] / M.N_SIMS), f_in, '0.0%'); put(ws, (h0 + 2 + i, 3), float(cbc[i] / M.N_SIMS), f_in, '0.0%')
    ch = BarChart(); ch.type = "col"; ch.title = "Unilever: NPV benefit distribution"; ch.height = 8; ch.width = 18
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=h0 + 1, max_row=h0 + 21), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=h0 + 2, max_row=h0 + 21)); ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.y_axis.number_format = '0%'; ws.add_chart(ch, f"E{h0+1}")
    widths(ws, {"A": 22, "B": 12, "C": 24, "D": 11, "E": 11, "F": 11, "G": 11, "H": 12, "I": 16, "J": 18})
    ws.freeze_panes = "A5"
    return ws

def build_sens(wb):
    ws = wb.create_sheet("Sensitivity")
    title(ws, "Sensitivity & simulation grids (STATIC values from the Python model)", "One-at-a-time tornado, and the traceability-coverage × supplier-substitution grid required by the project brief.")
    base, tor = M.tornado("ULVR", "BC")
    put(ws, "A4", "Unilever, permissioned blockchain — base-case NPV benefit (USD m)", f_bold); put(ws, "D4", float(base), f_in, NUM, fill=fill_key)
    hdr(ws, 5, ["Driver (low case / high case)", "Δ NPV, low case", "Δ NPV, high case", "Swing"], height=30)
    for i, x in tor.iterrows():
        r = 6 + i
        put(ws, (r, 1), x.driver); put(ws, (r, 2), float(x.low_case), f_in, NUM); put(ws, (r, 3), float(x.high_case), f_in, NUM); put(ws, (r, 4), f"=ABS(C{r}-B{r})", f_calc, NUM)
    n = len(tor)
    ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.overlap = 100; ch.title = "Tornado: change in NPV vs base (USD m)"; ch.height = 9; ch.width = 20
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=5, max_row=5 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + n)); ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.scaling.orientation = "maxMin"; ws.add_chart(ch, f"F4")
    note(ws, f"A{7+n}", "Regulatory probability and the realised fine rate enter the model multiplicatively, so they have identical swings — treat them as one 'expected regulatory loss' driver. "
         "The result is most sensitive to how likely and how costly enforcement is, and to programme cost and ramp-up speed — not to the blockchain-specific effectiveness number.")
    ws.merge_cells(start_row=7 + n, start_column=1, end_row=7 + n, end_column=4); ws.row_dimensions[7 + n].height = 78
    r = 10 + n + 3
    for pim, lab in ((1.0, "base substitution premium"), (2.5, "severe substitution premium (×2.5)")):
        sec(ws, r, f"Unilever — PV of supplier-substitution cost (USD m) by steady traceable share (rows) and share of non-traceable volume excluded (columns); {lab}", 6); r += 1
        g = M.coverage_exclusion_grid("ULVR", pim)
        hdr(ws, r, ["Traceable share ↓ / excluded →"] + list(g.columns)); r += 1
        r_first = r
        for idx, row in g.iterrows():
            put(ws, (r, 1), idx, f_bold)
            for j, v in enumerate(row): put(ws, (r, 2 + j), float(abs(v)), f_in, NUM)
            r += 1
        ws.conditional_formatting.add(f"B{r_first}:F{r-1}", ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"))
        r += 2
    note(ws, f"A{r}", "Risk heat map: green = cheap to substitute (high coverage or few exclusions), red = costly. Moving traceable share from 30% to 80% removes ~70% of the substitution bill at any exclusion rate — that value comes from traceability itself, whichever technology delivers it.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 44
    widths(ws, {"A": 46, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})
    return ws
